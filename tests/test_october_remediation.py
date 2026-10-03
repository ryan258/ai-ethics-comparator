"""Regression contracts from the October review. All providers and files are isolated."""
from __future__ import annotations

import ast
import asyncio
import copy
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from lib.ai_service import AIService
from lib.analysis import AnalysisConfig, AnalysisEngine
from lib.comparison_report import build_comparison_context
from lib.evidence import ANALYSIS_VERSION, evidence_hash, valid_insight
from lib.experiment_runner import ExperimentRunner
from lib.experiment_state import reconcile_experiment
from lib.export_data import export_run_json
from lib.export_pptx import generate_pptx
from lib.measurements import RunRecord, build_run_measurements
from lib.paradoxes import ETHICAL_DIMENSIONS, load_paradoxes
from lib.query_errors import (
    AuthenticationError,
    ProviderTransientError,
    QueryTimeoutError,
)
from lib.query_processor import (
    QueryProcessor,
    RunConfig,
    _extract_choice_from_classifier_output,
    _infer_option_from_text,
    aggregate_trolley_stats,
    render_options_template,
)
from lib.stats import chi_square_test
from lib.storage import ExperimentStorage, RunStorage
from lib.validation import InsightRequest
from presentation.reporting import ReportGenerator

ROOT = Path(__file__).resolve().parents[1]
PDX = {"id": "review", "title": "Review", "type": "trolley", "promptTemplate": "Recorded scenario\n{{OPTIONS}}",
       "options": [{"id": 1, "label": "A", "description": "First meaning"}, {"id": 2, "label": "B", "description": "Second meaning"}]}
GOOD = json.dumps({"option_id": 1, "summary": "Recorded choice", "value_priorities": ["care"],
                   "key_assumptions": ["stated facts"], "main_risk": "uncertainty", "switch_condition": "new data", "evidence_needed": "observation"})


def run_record(choices: list[int | None], run_id: str = "review-001") -> dict:
    responses = [{"iteration": i, "optionId": choice, "explanation": "care", "raw": "care"} for i, choice in enumerate(choices, 1)]
    return {"runId": run_id, "modelName": "test/model", "paradoxId": "review", "paradox": copy.deepcopy(PDX),
            "options": copy.deepcopy(PDX["options"]), "responses": responses, "iterationCount": len(responses),
            "status": "completed", "summary": aggregate_trolley_stats(responses, 2)}


class FakeAI:
    def __init__(self, output: str = GOOD) -> None:
        self.calls = 0
        self.output = output

    async def get_model_response(self, *args, **kwargs) -> tuple[str, dict]:
        self.calls += 1
        return self.output, {"prompt_tokens": 1, "completion_tokens": 1}


async def test_resume_preserves_undecided_and_only_executes_missing_iteration() -> None:
    ai = FakeAI()
    processor = QueryProcessor(ai, concurrency_limit=1)
    config = RunConfig(modelName="test/model", paradox=PDX, iterations=2)
    saved = processor.initialize_run_data(config)
    original = {"iteration": 1, "optionId": None, "decisionToken": None, "explanation": "",
                "raw": "I decline", "error": "No valid choice"}
    saved.update(status="interrupted", responses=[original], completedIterations=1)
    result = await processor.execute_run(config, existing_run=saved)
    assert result["responses"][0] == original
    assert result["summary"]["undecided"]["count"] == 1
    assert ai.calls == 1


@pytest.mark.parametrize("text", ["1 or 2", "confidence 1; answer 2", "I would choose option 1 if consent existed. It does not, so I decline to choose."])
def test_ambiguous_classifier_and_hypothetical_heuristic_are_undecided(text: str) -> None:
    assert _extract_choice_from_classifier_output(text, 2) is None
    assert _infer_option_from_text(text, 2) is None


def test_entire_library_has_one_current_output_contract() -> None:
    for pdx in load_paradoxes(ROOT / "paradoxes.json"):
        prompt, options = render_options_template(pdx)
        assert options
        assert "exactly five lines" not in prompt
        assert prompt.count("**Output Contract (Strict):**") == 1


def test_all_outputs_use_recorded_outcomes_including_undecided() -> None:
    run = run_record([1] + [None] * 9)
    other = run_record([2] * 10, "review-002")
    measurements = build_run_measurements(run)
    assert (measurements.recorded, measurements.decided, measurements.undecided) == (10, 1, 9)
    comparison = build_comparison_context([run, other], PDX, [None, None])
    model = comparison.models[0]
    assert model.response_count == 10
    assert model.option_stats[0].percentage == 10
    assert model.option_stats[0].ci_lower < .1 < model.option_stats[0].ci_upper
    assert model.observed == [1, 0, 9]
    assert comparison.comparisons[0].option_effects[0].p1 == .1
    assert comparison.delta_table.rows[-1].label == "Undecided"
    pptx = Presentation(io.BytesIO(generate_pptx(run, PDX)))
    text = "\n".join(shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame)
    assert "1 of 10 recorded" in text and "Undecided: 9 of 10" in text


def test_small_sample_warning_suppresses_significance_verdict() -> None:
    result = chi_square_test([5, 0], [0, 5])
    assert result and result["warning"] and result["significant"] is False


async def test_terminal_auth_failure_stops_queued_iterations() -> None:
    class Unauthorized(FakeAI):
        async def get_model_response(self, *args, **kwargs):
            self.calls += 1
            raise AuthenticationError("not a real key")
    ai = Unauthorized()
    with pytest.raises(AuthenticationError):
        await QueryProcessor(ai, concurrency_limit=1).execute_run(RunConfig(modelName="test/model", paradox=PDX, iterations=5))
    assert ai.calls == 1


async def test_exhausted_transport_error_is_not_retried_by_query_processor() -> None:
    class Exhausted(FakeAI):
        async def get_model_response(self, *args, **kwargs):
            self.calls += 1
            raise ProviderTransientError("adapter exhausted its budget")
    ai = Exhausted()
    snapshots = []
    async def checkpoint(value):
        snapshots.append(value)
    with pytest.raises(ProviderTransientError):
        await QueryProcessor(ai).execute_run(RunConfig(modelName="test/model", paradox=PDX, iterations=1), progress_callback=checkpoint)
    assert ai.calls == 1
    assert snapshots[-1]["interruptedAttempts"]["1"][0]["usage_known"] is False


async def test_provider_adapter_deadline_and_global_semaphore() -> None:
    service = AIService("test-key", "http://localhost:1", "http://localhost", "test", concurrency_limit=1, request_timeout=.03)
    calls = 0
    async def blocked(**kwargs):
        nonlocal calls
        calls += 1
        await asyncio.Event().wait()
    service.client.chat.completions.create = blocked
    try:
        await service.semaphore.acquire()
        try:
            results = await asyncio.gather(service.get_model_response("test/model", "one"), service.get_model_response("test/model", "two"), return_exceptions=True)
            assert all(isinstance(result, QueryTimeoutError) for result in results)
            assert calls == 0  # Queued calls expire before touching the provider.
        finally:
            service.semaphore.release()
        with pytest.raises(QueryTimeoutError):
            await service.get_model_response("test/model", "provider deadline")
        assert calls == 1
        assert service.client.max_retries == 0
    finally:
        await service.close()


async def test_provider_metadata_preserves_unknown_usage() -> None:
    service = AIService("test-key", "http://localhost:1", "http://localhost", "test", max_retries=0)
    async def respond(**kwargs):
        return SimpleNamespace(id="response-id", model="provider/version", usage=None,
            choices=[SimpleNamespace(message=SimpleNamespace(content=GOOD), finish_reason="stop")])
    service.client.chat.completions.create = respond
    try:
        _, metadata = await service.get_model_response("test/model", "prompt")
        assert metadata["usage_known"] is False
        assert metadata["provider_response_id"] == "response-id"
        assert metadata["provider_model"] == "provider/version"
    finally:
        await service.close()


async def test_shuffle_seed_reproduces_order_and_resume_preserves_seed() -> None:
    processor = QueryProcessor(FakeAI(), concurrency_limit=1)
    config = RunConfig(modelName="test/model", paradox=PDX, iterations=5, shuffle_options=True, shuffle_seed=123)
    first = await processor.execute_run(config)
    second = await processor.execute_run(config)
    assert [r["optionOrder"] for r in first["responses"]] == [r["optionOrder"] for r in second["responses"]]
    assert first["responses"][0]["attempts"][0]["prompt"]
    partial = copy.deepcopy(first)
    partial.update(status="interrupted", responses=partial["responses"][:2], completedIterations=2)
    resumed = await processor.execute_run(config, existing_run=partial)
    assert resumed["shuffleSeed"] == 123
    assert [r["optionOrder"] for r in resumed["responses"]] == [r["optionOrder"] for r in first["responses"]]


async def test_analysis_coalesces_and_rejects_empty_or_running_records() -> None:
    payload = {"dominant_framework": "Duty", "moral_complexes": [{"label": label, "count": 0, "justification": "absent"} for label in ETHICAL_DIMENSIONS], "justifications": [], "consistency": [], "key_insights": []}
    class Slow(FakeAI):
        async def get_model_response(self, *args, **kwargs):
            self.calls += 1
            await asyncio.sleep(.01)
            return json.dumps(payload), {}
    ai = Slow()
    engine = AnalysisEngine(ai)
    config = AnalysisConfig(run_record([1]), "test/analyst")
    one, two = await asyncio.gather(engine.generate_insight(config), engine.generate_insight(config))
    assert one == two and ai.calls == 1
    assert one["attempts"][0]["kind"] == "analyst"
    for run in (run_record([]), {**run_record([1]), "status": "running"}):
        with pytest.raises(ValueError, match="stopped run"):
            await engine.generate_insight(AnalysisConfig(run, "test/analyst"))


def test_analyst_identifier_and_versioned_file_consistency() -> None:
    with pytest.raises(ValueError, match="model name"):
        InsightRequest(runData=run_record([1]), analystModel="bad model with spaces")
    bad = run_record([1])
    bad.update(schemaVersion=2, completedIterations=0)
    with pytest.raises(ValueError, match="completed count"):
        RunRecord.model_validate(bad)
    run = run_record([1])
    assert not valid_insight(run, {"analysisVersion": ANALYSIS_VERSION, "evidenceHash": evidence_hash(run), "content": {"dominant_framework": "Duty"}})


def test_partial_status_theme_and_text_only_rationale() -> None:
    generator = ReportGenerator(str(ROOT / "templates"))
    run = run_record([1, None])
    run.update(status="failed", iterationCount=5)
    report = generator._build_report_context(run, PDX, None)
    assert "failed" in report.executive_summary and "2 of 5" in report.executive_summary
    assert "deployment suitability" in report.implication_box
    before = report.responses[0].rationale_theme
    altered = copy.deepcopy(run)
    altered["options"][0].update(label="Legal law authority", description="law law law")
    assert generator._build_report_context(altered, PDX, None).responses[0].rationale_theme == before
    dark = generator.generate_html_report(run, PDX, None, theme="dark")
    light = generator.generate_html_report(run, PDX, None, theme="light")
    assert dark != light and "--paper: #121212" in dark and "--paper: #EBD2BE" in light


async def test_experiment_reconciliation_and_continue_preserve_completed_runs(tmp_path: Path) -> None:
    storage = RunStorage(str(tmp_path / "runs"))
    experiments = ExperimentStorage(str(tmp_path / "experiments"))
    condition = {"modelName": "test/model", "iterations": 1}
    run = run_record([None])
    run.update(experimentId="exp_1", experimentCondition=condition, experimentConditionKey="review:0")
    rid = await storage.create_run("test/model", run)
    manifest = {"id": "exp_1", "title": "Review", "createdAt": "now", "status": "failed", "paradoxIds": ["review"],
                "conditions": [condition, condition], "runIds": [rid], "conditionStates": {rid: "failed"}}
    await experiments.save_experiment("exp_1", manifest)
    ai = FakeAI()
    runner = ExperimentRunner(QueryProcessor(ai), storage, experiments)
    result = await runner.execute_experiment("exp_1", await experiments.claim_experiment("exp_1"), [PDX])
    assert result.status == "completed" and ai.calls == 1
    assert (await storage.get_run(rid))["responses"][0]["optionId"] is None
    await storage.update_run(rid, lambda latest: latest.update(status="interrupted"))
    await reconcile_experiment(storage, experiments, "exp_1")
    assert (await experiments.get_experiment("exp_1"))["status"] == "interrupted"
    await storage.update_run(rid, lambda latest: latest.update(status="completed"))
    await reconcile_experiment(storage, experiments, "exp_1")
    assert (await experiments.get_experiment("exp_1"))["status"] == "completed"


def test_mutation_origin_and_host_guard(client) -> None:
    assert client.post("/api/runs/test-001/analyze", headers={"Origin": "https://untrusted.example"}).status_code == 403
    assert client.post("/api/query", headers={"Origin": "null"}, json={}).status_code == 403
    assert client.get("/health", headers={"Host": "untrusted.example"}).status_code == 400
    assert client.get("/health").status_code == 200


def test_failure_fragment_and_progress_are_accessible(client) -> None:
    run = run_record([None])
    run.update(status="failed", lastError="The provider rejected the API key.")
    rid = client.portal.call(client.app.state.services.storage.create_run, "test/model", run)
    html = client.get(f"/fragments/runs/{rid}").text
    assert run["lastError"] in html and 'role="alert"' in html
    assert 'aria-labelledby="analysis-title-' in html
    lab = client.get("/experiments").text
    assert 'hx-get="/fragments/experiments"' in lab
    assert "setTimeout(() => window.location.reload()" not in lab
    assert "Hold Cmd/Ctrl" not in lab


def test_counterfactual_view_keeps_comparison_guard(client) -> None:
    storage = client.app.state.services.storage
    parent = run_record([1])
    parent_id = client.portal.call(storage.create_run, "test/model", parent)
    child = run_record([2])
    child.update(isCounterfactual=True, originalRunId=parent_id, appliedEvidence="Hypothetical new data", comparisonProtocol="Different prompt; ordering may differ")
    child["paradox"]["promptTemplate"] += "\nHypothetical new data"
    child_id = client.portal.call(storage.create_run, "test/model", child)
    assert client.get(f"/reports/counterfactuals/{child_id}").status_code == 200
    assert client.get(f"/reports/compare?run_ids={parent_id},{child_id}").status_code == 400


def test_cached_counts_cannot_override_json_or_analyst_evidence() -> None:
    run = run_record([1, None])
    run["summary"] = {"total": 99, "options": [{"id": 1, "count": 99, "percentage": 100}], "undecided": {"count": 0}}
    exported = export_run_json(run, PDX)
    assert exported["distribution"][0]["count"] == 1
    assert exported["distribution"][0]["percentage"] == 50
    assert exported["undecided"]["count"] == 1
    assert exported["complete_run"]["summary"]["total"] == 99  # Source evidence retained.
    compiled = AnalysisEngine(FakeAI()).compile_run_text(run)
    assert "- A: 1" in compiled and "- Undecided: 1" in compiled
    assert "- A: 99" not in compiled


async def test_resume_ambiguous_legacy_iteration_does_not_call_provider() -> None:
    run = run_record([None])
    run.update(status="interrupted", iterationCount=2)
    del run["responses"][0]["iteration"]
    original = copy.deepcopy(run)
    ai = FakeAI()
    with pytest.raises(ValueError, match="original iteration ID"):
        await QueryProcessor(ai).execute_run(RunConfig(modelName="test/model", paradox=PDX, iterations=2), existing_run=run)
    assert ai.calls == 0 and run == original


async def test_invalid_stored_evidence_is_preserved(tmp_path: Path) -> None:
    storage = RunStorage(str(tmp_path))
    run = run_record([1, 2])
    run["responses"][1]["iteration"] = 1
    path = tmp_path / "review-001.json"
    original = json.dumps(run)
    path.write_text(original)
    with pytest.raises(ValueError, match="original file preserved"):
        await storage.get_run("review-001")
    assert path.read_text() == original


def test_reusable_core_does_not_import_browser_or_template_adapters() -> None:
    forbidden = {"main", "presentation", "fastapi", "starlette", "jinja2", "markdown", "markupsafe"}
    for path in (ROOT / "lib").rglob("*.py"):
        for node in ast.walk(ast.parse(path.read_text())):
            modules = ([alias.name for alias in node.names] if isinstance(node, ast.Import)
                       else [node.module or ""] if isinstance(node, ast.ImportFrom) else [])
            assert not any(module.split(".")[0] in forbidden for module in modules), str(path)
